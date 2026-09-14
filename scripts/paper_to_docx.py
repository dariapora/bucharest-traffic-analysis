import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "paper" / "recurrent_congestion_bucharest.md"
DST = SRC.with_suffix(".docx")

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_runs(par, text, base_bold=False, base_italic=False, size=None):
    for piece in INLINE.split(text):
        if not piece:
            continue
        bold, italic, mono = base_bold, base_italic, False
        if piece.startswith("**") and piece.endswith("**"):
            piece, bold = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*"):
            piece, italic = piece[1:-1], True
        elif piece.startswith("`") and piece.endswith("`"):
            piece, mono = piece[1:-1], True
        run = par.add_run(piece)
        run.bold = bold
        run.italic = italic
        if mono:
            run.font.name = "Consolas"
            run.font.size = Pt((size or 11) - 1)
        elif size:
            run.font.size = Pt(size)


def set_cell_shading(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_table(doc, rows):
    header, body = rows[0], [r for r in rows[1:]]
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        add_runs(cell.paragraphs[0], text, base_bold=True, size=9)
        set_cell_shading(cell, "E7E6E6")
    for row in body:
        cells = table.add_row().cells
        for i in range(ncols):
            text = row[i] if i < len(row) else ""
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], text, size=9)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
    doc.add_paragraph()


def split_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def is_separator(line):
    return re.fullmatch(r"\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?", line.strip()) is not None


def build(doc, lines):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.5)

    i = 0
    n = len(lines)
    first_heading = True
    while i < n:
        line = lines[i].rstrip("\n")
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if first_heading and level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(p, text, base_bold=True, size=16)
                first_heading = False
            else:
                h = doc.add_heading(level=min(level, 3))
                add_runs(h, text)
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                if not is_separator(lines[i]):
                    rows.append(split_row(lines[i]))
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        if stripped.startswith(">"):
            block = []
            while i < n and lines[i].strip().startswith(">"):
                block.append(lines[i].strip()[1:].strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.right_indent = Cm(1)
            add_runs(p, " ".join(b for b in block if b), base_italic=True, size=10)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            continue

        if stripped.startswith("- "):
            while i < n and lines[i].strip().startswith("- "):
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, lines[i].strip()[2:])
                i += 1
            continue

        if re.match(r"^(Authors|Affiliation):", stripped):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, stripped)
            i += 1
            continue

        # paragraph: join consecutive non-blank, non-special lines
        block = [stripped]
        i += 1
        while i < n:
            nxt = lines[i].strip()
            if not nxt or nxt.startswith(("#", "|", ">", "- ", "---")):
                break
            block.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, " ".join(block))


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    build(doc, lines)
    doc.save(DST)
    print(f"written {DST}")


if __name__ == "__main__":
    main()
